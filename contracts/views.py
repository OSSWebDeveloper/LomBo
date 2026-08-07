# -*- coding: utf-8 -*-
import json
import os
import re
import shutil

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Count, F, Q, Sum
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from accounts.models import User

from docx import Document

from .docgen import build_contract_docx, build_garov_docx, payment_schedule
from .pdf import PdfImkoniYoq, docx_dan_pdf, pdf_imkoni_bor
from .shablondan import SHABLONLAR, garovi_bormi
from .forms import (ContractForm, ContractSearchForm, DeleteRequestForm, GuarantorForm,
                    JewelryFormSet, VehicleForm)
from .models import Amal, Contract, DeleteRequest, amal_yoz


# --------------------------------------------------------------- yordamchi

def _boss_required(user):
    if not user.is_boshliq:
        raise PermissionDenied('Bu bo\'lim faqat boshliqlar uchun.')


def _can_view_contract(user, contract):
    """Ishchi faqat o'zinikini, boshliq hammasini ko'radi."""
    return user.is_boshliq or contract.created_by_id == user.id


def _tahrir_huquqi(user, contract):
    """Kim tahrirlay oladi.

    Vakolatli boshliq — istalgan shartnomani, istalgan vaqtda.
    Ishchi — faqat o'zi endigina kiritgan shartnomani (ISHCHI_TAHRIR_DAQIQA).
    """
    return bool(user.is_superuser
                or (user.is_boshliq and user.can_full_edit)
                or contract.ishchi_tahrirlay_oladi(user))


def _tahrir_rad_sababi(user, contract):
    """Tahrirlab bo'lmasa — foydalanuvchiga nima qilishini aytadigan izoh."""
    if user.is_boshliq:
        return ("Shartnomani to'liq tahrirlash uchun maxsus vakolat kerak. "
                "Vakolatni administrator admin panel orqali beradi.")
    if contract.created_by_id != user.id:
        return 'Bu shartnoma sizga tegishli emas.'
    if contract.status != Contract.STATUS_ACTIVE:
        return ("Bu shartnoma uchun o'chirish so'rovi yuborilgan — "
                "boshliq javob bergunicha uni o'zgartirib bo'lmaydi.")
    daqiqa = getattr(settings, 'ISHCHI_TAHRIR_DAQIQA', 0)
    if not daqiqa:
        return ("Shartnomani tahrirlash ishchilar uchun yopilgan. "
                "O'zgartirish uchun boshlig'ingizga murojaat qiling.")
    return (f"Shartnomani faqat kiritilgandan keyingi {daqiqa} daqiqa ichida "
            f"tuzatish mumkin, bu muddat o'tib ketdi. "
            f"O'zgartirish uchun boshlig'ingizga murojaat qiling.")


def _garov_bahosini_yangila(contract):
    """Zargarlikda garov bahosi jadvaldagi summalar yig'indisiga teng."""
    if contract.collateral_type != Contract.TYPE_ZARGARLIK:
        return
    jami = contract.jewelry_items.aggregate(s=Sum('value'))['s'] or 0
    if contract.garov_value != jami:
        contract.garov_value = jami
        contract.save(update_fields=['garov_value'])


# --------------------------------------------------------------- dashboard

@login_required
def dashboard(request):
    user = request.user
    if user.is_boshliq:
        workers = User.objects.filter(role=User.ROLE_ISHCHI)
        if not user.is_superuser:
            workers = workers.filter(added_by=user)
        contracts = Contract.objects.all()
        pending = DeleteRequest.objects.filter(status=DeleteRequest.STATUS_PENDING)
        if not user.is_superuser:
            pending = pending.filter(assigned_to=user)
        ctx = {
            'is_boss': True,
            'workers_count': workers.count(),
            'contracts_count': contracts.count(),
            'total_amount': contracts.aggregate(s=Sum('amount'))['s'] or 0,
            'pending_count': pending.count(),
            'pending_requests': pending.select_related('requested_by')[:5],
            'recent': contracts.select_related('created_by')[:10],
            'top_workers': workers.annotate(
                live=Count('contracts')).order_by('-contracts_added')[:5],
        }
    else:
        my = Contract.objects.filter(created_by=user)
        ctx = {
            'is_boss': False,
            'my_count': my.count(),
            'total_added': user.contracts_added,
            'my_amount': my.aggregate(s=Sum('amount'))['s'] or 0,
            'recent': my[:10],
            'my_requests': DeleteRequest.objects.filter(requested_by=user)[:5],
        }
    return render(request, 'contracts/dashboard.html', ctx)


# --------------------------------------------------------------- ro'yxat

@login_required
def contract_list(request):
    user = request.user
    qs = Contract.objects.select_related('created_by')
    if not user.is_boshliq:
        qs = qs.filter(created_by=user)

    users = None
    if user.is_boshliq:
        users = User.objects.all().order_by('username')
    form = ContractSearchForm(request.GET or None, users=users)
    if form.is_valid():
        q = form.cleaned_data.get('q')
        if q:
            cond = Q(borrower_fio__icontains=q)
            if q.strip().isdigit():
                cond |= Q(number=int(q.strip()))
            qs = qs.filter(cond)
        if form.cleaned_data.get('collateral_type'):
            qs = qs.filter(collateral_type=form.cleaned_data['collateral_type'])
        if form.cleaned_data.get('date_from'):
            qs = qs.filter(date__gte=form.cleaned_data['date_from'])
        if form.cleaned_data.get('date_to'):
            qs = qs.filter(date__lte=form.cleaned_data['date_to'])
        if form.cleaned_data.get('created_by'):
            qs = qs.filter(created_by_id=form.cleaned_data['created_by'])

    return render(request, 'contracts/contract_list.html', {
        'contracts': qs, 'form': form, 'total': qs.count(),
    })


# --------------------------------------------------------------- qo'shish / tahrirlash

def _collateral_forms(request, contract=None, data=None):
    """Ta'minot turiga qarab qo'shimcha formalar."""
    jewelry = JewelryFormSet(data, instance=contract, prefix='jewelry')
    vehicle = VehicleForm(data, instance=getattr(contract, 'vehicle', None), prefix='vehicle')
    guarantor = GuarantorForm(data, instance=getattr(contract, 'guarantor', None),
                              prefix='guarantor')
    return jewelry, vehicle, guarantor


@login_required
def contract_create(request):
    data = request.POST or None
    form = ContractForm(data)   # raqamni forma o'zi avtomatik beradi
    jewelry, vehicle, guarantor = _collateral_forms(request, data=data)

    if request.method == 'POST':
        tur = request.POST.get('collateral_type')
        ok = form.is_valid()
        extra_ok = True
        if tur == Contract.TYPE_ZARGARLIK:
            extra_ok = jewelry.is_valid()
        elif tur == Contract.TYPE_TRANSPORT:
            extra_ok = vehicle.is_valid()
        elif tur == Contract.TYPE_KAFILLIK:
            extra_ok = guarantor.is_valid()

        if ok and extra_ok:
            contract = form.save(commit=False)
            contract.end_date = form.cleaned_data['end_date']
            contract.created_by = request.user
            contract.save()

            if tur == Contract.TYPE_ZARGARLIK:
                jewelry.instance = contract
                items = jewelry.save()
                _garov_bahosini_yangila(contract)
                if not items:
                    messages.warning(request, 'Zargarlik jadvali bo\'sh qoldi.')
            elif tur == Contract.TYPE_TRANSPORT:
                v = vehicle.save(commit=False)
                v.contract = contract
                v.save()
            elif tur == Contract.TYPE_KAFILLIK:
                g = guarantor.save(commit=False)
                g.contract = contract
                g.save()

            User.objects.filter(pk=request.user.pk).update(
                contracts_added=F('contracts_added') + 1)
            amal_yoz(request.user, Amal.YARATDI, f'Shartnoma №{contract.number}',
                     f'{contract.get_collateral_type_display()}, {contract.borrower_fio}')
            messages.success(request, f'Shartnoma №{contract.number} qo\'shildi.')
            return redirect('contract_detail', pk=contract.pk)
        messages.error(request, 'Formada xatolar bor, tekshirib chiqing.')

    return render(request, 'contracts/contract_form.html', {
        'form': form, 'jewelry': jewelry, 'vehicle': vehicle, 'guarantor': guarantor,
        'title': "Yangi shartnoma qo'shish", 'is_edit': False,
    })


@login_required
def contract_edit(request, pk):
    """Tahrirlash — vakolatli boshliq yoki shartnomani endigina kiritgan ishchi."""
    contract = get_object_or_404(Contract, pk=pk)
    if not _tahrir_huquqi(request.user, contract):
        raise PermissionDenied(_tahrir_rad_sababi(request.user, contract))
    # Ishchining tuzatishi: muddat GET va POST'da alohida tekshiriladi,
    # ya'ni forma ochiq turganda muddat tugasa saqlab bo'lmaydi.
    ishchi_tahriri = request.user.is_ishchi

    data = request.POST or None
    form = ContractForm(data, instance=contract)
    jewelry, vehicle, guarantor = _collateral_forms(request, contract=contract, data=data)

    if request.method == 'POST':
        tur = request.POST.get('collateral_type')
        ok = form.is_valid()
        extra_ok = True
        if tur == Contract.TYPE_ZARGARLIK:
            extra_ok = jewelry.is_valid()
        elif tur == Contract.TYPE_TRANSPORT:
            extra_ok = vehicle.is_valid()
        elif tur == Contract.TYPE_KAFILLIK:
            extra_ok = guarantor.is_valid()

        if ok and extra_ok:
            contract = form.save(commit=False)
            contract.end_date = form.cleaned_data['end_date']
            contract.save()

            # Ta'minot turi o'zgargan bo'lsa, eskisining ma'lumotlarini tozalash
            if tur != Contract.TYPE_ZARGARLIK:
                contract.jewelry_items.all().delete()
            if tur != Contract.TYPE_TRANSPORT and hasattr(contract, 'vehicle'):
                contract.vehicle.delete()
            if tur != Contract.TYPE_KAFILLIK and hasattr(contract, 'guarantor'):
                contract.guarantor.delete()

            if tur == Contract.TYPE_ZARGARLIK:
                jewelry.instance = contract
                jewelry.save()
                _garov_bahosini_yangila(contract)
            elif tur == Contract.TYPE_TRANSPORT:
                v = vehicle.save(commit=False)
                v.contract = contract
                v.save()
            elif tur == Contract.TYPE_KAFILLIK:
                g = guarantor.save(commit=False)
                g.contract = contract
                g.save()

            # Logda faqat o'zgargan qismlar soni saqlanadi
            soni = len(form.changed_data)
            if tur == Contract.TYPE_ZARGARLIK:
                soni += sum(len(f.changed_data) for f in jewelry.forms)
            elif tur == Contract.TYPE_TRANSPORT:
                soni += len(vehicle.changed_data)
            elif tur == Contract.TYPE_KAFILLIK:
                soni += len(guarantor.changed_data)
            izoh = f'{soni} ta qism o\'zgartirildi'
            if ishchi_tahriri:
                izoh += ' (kiritgandan keyingi tuzatish)'
            amal_yoz(request.user, Amal.OZGARTIRDI, f'Shartnoma №{contract.number}', izoh)
            messages.success(request, f'Shartnoma №{contract.number} o\'zgartirildi.')
            return redirect('contract_detail', pk=contract.pk)
        messages.error(request, 'Formada xatolar bor.')

    return render(request, 'contracts/contract_form.html', {
        'form': form, 'jewelry': jewelry, 'vehicle': vehicle, 'guarantor': guarantor,
        'title': (f"Shartnoma №{contract.number} — tuzatish" if ishchi_tahriri
                  else f"Shartnoma №{contract.number} — to'liq tahrirlash"),
        'is_edit': True, 'contract': contract,
        'tahrir_qoldiq': contract.tahrir_qoldiq_daqiqa if ishchi_tahriri else 0,
    })


@login_required
def contract_detail(request, pk):
    contract = get_object_or_404(
        Contract.objects.select_related('created_by'), pk=pk)
    if not _can_view_contract(request.user, contract):
        raise PermissionDenied('Bu shartnoma sizga tegishli emas.')
    can_edit = _tahrir_huquqi(request.user, contract)
    ishchi_tuzatishi = contract.ishchi_tahrirlay_oladi(request.user)
    pending = contract.delete_requests.filter(
        status=DeleteRequest.STATUS_PENDING).first()
    return render(request, 'contracts/contract_detail.html', {
        'contract': contract, 'can_edit': can_edit,
        # O'chirish tugmasi faqat boshliqda: ishchining o'chirish so'rovi
        # oqimi hozircha saytda ochilmagan (qarang: contract_delete).
        'can_delete': request.user.is_boshliq,
        'tahrir_qoldiq': contract.tahrir_qoldiq_daqiqa if ishchi_tuzatishi else 0,
        'schedule': payment_schedule(contract),
        'pending_request': pending,
        'pdf_bor': pdf_imkoni_bor(),
        'garov_bor': garovi_bormi(contract),
    })


@login_required
def contract_download_pdf(request, pk):
    """Shartnomani PDF ko'rinishida yuklab olish."""
    contract = get_object_or_404(Contract, pk=pk)
    if not _can_view_contract(request.user, contract):
        raise PermissionDenied('Bu shartnoma sizga tegishli emas.')
    try:
        data = docx_dan_pdf(build_contract_docx(contract))
    except PdfImkoniYoq as xato:
        messages.error(request, str(xato))
        return redirect('contract_detail', pk=pk)
    resp = HttpResponse(data, content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="shartnoma_{contract.number}.pdf"'
    return resp


@login_required
def contract_download(request, pk):
    contract = get_object_or_404(Contract, pk=pk)
    if not _can_view_contract(request.user, contract):
        raise PermissionDenied('Bu shartnoma sizga tegishli emas.')
    data = build_contract_docx(contract)
    resp = HttpResponse(
        data, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    resp['Content-Disposition'] = f'attachment; filename="shartnoma_{contract.number}.docx"'
    return resp


# --------------------------------------------------------------- o'chirish oqimi

@login_required
def contract_delete(request, pk):
    """Ishchi -> so'rov yuboradi. Boshliq -> to'g'ridan-to'g'ri o'chiradi."""
    contract = get_object_or_404(Contract, pk=pk)
    if not _can_view_contract(request.user, contract):
        raise PermissionDenied('Bu shartnoma sizga tegishli emas.')

    if request.user.is_boshliq:
        if request.method == 'POST':
            num = contract.number
            fio = contract.borrower_fio
            contract.delete()
            amal_yoz(request.user, Amal.OCHIRDI, f'Shartnoma №{num}', fio)
            messages.success(request, f'Shartnoma №{num} o\'chirildi.')
            return redirect('contract_list')
        return render(request, 'contracts/contract_confirm_delete.html', {
            'contract': contract, 'direct': True})

    # Ishchi — so'rov
    existing = contract.delete_requests.filter(
        status=DeleteRequest.STATUS_PENDING).first()
    if existing:
        messages.info(request, 'Bu shartnoma uchun so\'rov allaqachon yuborilgan.')
        return redirect('contract_detail', pk=pk)

    boss = request.user.added_by
    if boss is None:
        messages.error(
            request, "Sizni qo'shgan boshliq topilmadi. Administratorga murojaat qiling.")
        return redirect('contract_detail', pk=pk)

    form = DeleteRequestForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        DeleteRequest.objects.create(
            contract=contract,
            contract_number=contract.number,
            contract_info=f'{contract.borrower_fio} — {contract.amount} so\'m',
            requested_by=request.user,
            assigned_to=boss,
            reason=form.cleaned_data['reason'],
        )
        contract.status = Contract.STATUS_PENDING_DELETE
        contract.save(update_fields=['status'])
        amal_yoz(request.user, Amal.SOROV_YUBORDI, f'Shartnoma №{contract.number}',
                 f'{boss} ga yuborildi')
        messages.success(
            request, f"O'chirish so'rovi {boss} ga yuborildi. Tasdiqlansa o'chiriladi.")
        return redirect('contract_detail', pk=pk)

    return render(request, 'contracts/contract_confirm_delete.html', {
        'contract': contract, 'direct': False, 'form': form, 'boss': boss})


@login_required
def delete_requests(request):
    """Boshliq uchun so'rovlar ro'yxati."""
    _boss_required(request.user)
    qs = DeleteRequest.objects.select_related('requested_by', 'contract')
    if not request.user.is_superuser:
        qs = qs.filter(assigned_to=request.user)
    return render(request, 'contracts/delete_requests.html', {
        'requests': qs,
        'pending_count': qs.filter(status=DeleteRequest.STATUS_PENDING).count(),
    })


@login_required
def delete_request_decide(request, pk, action):
    _boss_required(request.user)
    req = get_object_or_404(DeleteRequest, pk=pk)
    if not request.user.is_superuser and req.assigned_to_id != request.user.id:
        raise PermissionDenied('Bu so\'rov sizga yuborilmagan.')
    if req.status != DeleteRequest.STATUS_PENDING:
        messages.info(request, 'Bu so\'rov allaqachon ko\'rib chiqilgan.')
        return redirect('delete_requests')

    if request.method != 'POST':
        return redirect('delete_requests')

    if action == 'approve':
        contract = req.contract
        req.status = DeleteRequest.STATUS_APPROVED
        req.decided_at = timezone.now()
        req.save()
        if contract:
            num = contract.number
            contract.delete()   # DeleteRequest.contract = SET_NULL, so'rov tarixi qoladi
            amal_yoz(request.user, Amal.SOROV_TASDIQLADI, f'Shartnoma №{num}',
                     f"{req.requested_by} so'ragan, o'chirildi")
            messages.success(request, f"So'rov tasdiqlandi, shartnoma №{num} o'chirildi.")
        else:
            messages.info(request, 'Shartnoma allaqachon o\'chirilgan.')
    else:
        req.status = DeleteRequest.STATUS_REJECTED
        req.decided_at = timezone.now()
        req.save()
        if req.contract:
            req.contract.status = Contract.STATUS_ACTIVE
            req.contract.save(update_fields=['status'])
        amal_yoz(request.user, Amal.SOROV_RAD_ETDI,
                 f'Shartnoma №{req.contract_number}', f"{req.requested_by} so'ragan")
        messages.success(request, "So'rov rad etildi, shartnoma saqlanib qoldi.")
    return redirect('delete_requests')


# --------------------------------------------------------------- monitoring

@login_required
def monitoring(request):
    _boss_required(request.user)
    workers = User.objects.filter(role=User.ROLE_ISHCHI)
    if not request.user.is_superuser:
        workers = workers.filter(added_by=request.user)
    workers = workers.annotate(
        live_count=Count('contracts'),
        live_amount=Sum('contracts__amount'),
    ).order_by('-contracts_added')

    contracts = Contract.objects.all()
    by_type = (contracts.values('collateral_type')
               .annotate(c=Count('id'), s=Sum('amount')).order_by('-c'))
    type_labels = dict(Contract.TYPE_CHOICES)
    by_type = [{'label': type_labels.get(r['collateral_type'], r['collateral_type']),
                'count': r['c'], 'sum': r['s'] or 0} for r in by_type]

    by_month = (contracts.values('date__year', 'date__month')
                .annotate(c=Count('id'), s=Sum('amount'))
                .order_by('-date__year', '-date__month')[:12])

    return render(request, 'contracts/monitoring.html', {
        'workers': workers,
        'by_type': by_type,
        'by_month': by_month,
        'total_count': contracts.count(),
        'total_amount': contracts.aggregate(s=Sum('amount'))['s'] or 0,
        'pending_deletes': DeleteRequest.objects.filter(
            status=DeleteRequest.STATUS_PENDING).count(),
    })


# --------------------------------------------------------------- shablonlar

SHABLON_TURLARI = {
    'zargarlik': ('zargarlik.docx', 'Zargarlik buyumlari garovi',
                  ['{{ fio }}', '{{ summa }}', '{{ raqam }}', '#NOMI#']),
    'transport': ('transport.docx', 'Transport vositasi garovi',
                  ['{{ fio }}', '{{ summa }}', '{{ raqam }}', '{{ rusumi }}']),
    'kafillik': ('kafillik.docx', 'Ish haqi kafilligi',
                 ['{{ fio }}', '{{ summa }}', '{{ raqam }}', '{{ kafil }}']),
}


def _shablon_vakolati(user):
    if not user.shablon_vakolati:
        raise PermissionDenied(
            "Shablonni o'zgartirish uchun maxsus vakolat kerak. "
            "Vakolatni administrator admin panel orqali beradi.")


@login_required
def shablon_list(request):
    _shablon_vakolati(request.user)
    qatorlar = []
    for turi, (fayl, nom, _belgilar) in SHABLON_TURLARI.items():
        yol = os.path.join(SHABLONLAR, fayl)
        mavjud = os.path.exists(yol)
        qatorlar.append({
            'turi': turi, 'nom': nom, 'fayl': fayl, 'mavjud': mavjud,
            'hajmi': round(os.path.getsize(yol) / 1024) if mavjud else 0,
            'ozgartirilgan': (timezone.datetime.fromtimestamp(os.path.getmtime(yol))
                              if mavjud else None),
        })
    return render(request, 'contracts/shablon_list.html', {
        'shablonlar': qatorlar,
        'zaxiralar': _zaxiralar(),
    })


def _zaxiralar():
    papka = os.path.join(SHABLONLAR, 'zaxira')
    if not os.path.isdir(papka):
        return []
    fayllar = sorted(os.listdir(papka), reverse=True)[:10]
    return [{'nom': f,
             'vaqt': timezone.datetime.fromtimestamp(
                 os.path.getmtime(os.path.join(papka, f)))} for f in fayllar]


@login_required
def shablon_download(request, turi):
    _shablon_vakolati(request.user)
    if turi not in SHABLON_TURLARI:
        raise Http404
    fayl = SHABLON_TURLARI[turi][0]
    yol = os.path.join(SHABLONLAR, fayl)
    if not os.path.exists(yol):
        raise Http404
    with open(yol, 'rb') as f:
        resp = HttpResponse(
            f.read(),
            content_type='application/vnd.openxmlformats-officedocument'
                         '.wordprocessingml.document')
    resp['Content-Disposition'] = f'attachment; filename="{fayl}"'
    return resp


def _shablonni_tekshir(vaqtinchalik_yol, turi):
    """Yuklangan fayl haqiqiy shablonmi? (xato bo'lsa matn qaytaradi)"""
    kerakli = SHABLON_TURLARI[turi][2]
    try:
        doc = Document(vaqtinchalik_yol)
    except Exception:
        return 'Fayl Word hujjati emas yoki buzilgan.'

    matn = '\n'.join(p.text for p in _barcha_xatboshilar(doc))
    yoq = [b for b in kerakli if b.replace(' ', '') not in matn.replace(' ', '')]
    if yoq:
        return ('Shablonda kerakli belgilar topilmadi: ' + ', '.join(yoq) +
                '. Shablonni yuklab olib, faqat matnini tahrirlang — '
                'belgilarga tegmang.')

    # Sinov uchun to'ldirib ko'ramiz — Jinja xatolarini shu yerda tutamiz
    try:
        from docxtpl import DocxTemplate
        tpl = DocxTemplate(vaqtinchalik_yol)
        tpl.render(_sinov_konteksti())
    except Exception as xato:
        return f'Shablonni to\'ldirib ko\'rishda xato: {xato}'
    return None


def _barcha_xatboshilar(doc):
    from docx.oxml.ns import qn as _qn
    from docx.text.paragraph import Paragraph as _P
    for el in doc.element.body.iter(_qn('w:p')):
        yield _P(el, None)


def _sinov_konteksti():
    nomlar = ['raqam', 'garov_raqam', 'sana', 'sana_soz', 'tugash', 'muddat',
              'muddat_soz', 'foiz', 'foiz_soz', 'summa', 'fio', 'pasport', 'manzil',
              'telefon', 'daromad', 'pasport_seriya', 'pasport_soni', 'pasport_sana',
              'pasport_viloyat', 'pasport_bolim',
              'garov_baho', 'garov_baho_raqam', 'jami_soni', 'jami_ogirligi',
              'kafil', 'kafillik_summa', 'garov_mulki', 'garov_egasi',
              'garov_rahbari', 'garov_rahbari_qisqa', 'summa_raqam_soz',
              'davlat_raqami', 'rusumi', 'rangi', 'shassi', 'yili', 'texpasport']
    return {n: 'X' for n in nomlar}


@login_required
def shablon_upload(request, turi):
    _shablon_vakolati(request.user)
    if turi not in SHABLON_TURLARI:
        raise Http404
    fayl, nom, _belgilar = SHABLON_TURLARI[turi]

    if request.method != 'POST' or 'shablon' not in request.FILES:
        messages.error(request, 'Fayl tanlanmadi.')
        return redirect('shablon_list')

    yuklangan = request.FILES['shablon']
    if not yuklangan.name.lower().endswith('.docx'):
        messages.error(request, 'Faqat .docx fayl yuklanadi (.doc emas).')
        return redirect('shablon_list')

    vaqtinchalik = os.path.join(SHABLONLAR, f'.yangi_{fayl}')
    with open(vaqtinchalik, 'wb') as chiqish:
        for bolak in yuklangan.chunks():
            chiqish.write(bolak)

    xato = _shablonni_tekshir(vaqtinchalik, turi)
    if xato:
        os.remove(vaqtinchalik)
        messages.error(request, xato)
        return redirect('shablon_list')

    # Eskisini zaxiraga olamiz
    joriy = os.path.join(SHABLONLAR, fayl)
    if os.path.exists(joriy):
        zaxira_papka = os.path.join(SHABLONLAR, 'zaxira')
        os.makedirs(zaxira_papka, exist_ok=True)
        belgi = timezone.localtime().strftime('%Y%m%d-%H%M%S')
        shutil.copy2(joriy, os.path.join(zaxira_papka, f'{belgi}_{fayl}'))
        os.remove(joriy)
    os.replace(vaqtinchalik, joriy)

    amal_yoz(request.user, Amal.SHABLON_OZGARTIRDI, nom,
             f'fayl bilan almashtirildi ({yuklangan.name})')
    messages.success(request, f'«{nom}» shabloni almashtirildi. '
                              f'Eski nusxa zaxiraga saqlandi.')
    return redirect('shablon_list')


# --------------------------------------------------------------- onlayn tahrirlash

BELGI_NAQSHI = re.compile(r'(\{\{.*?\}\}|\{%.*?%\}|#[A-Z_]+#)')


def _shablon_xatboshilari(yol):
    """Shablondagi matnli xatboshilar: indeks, matn va rangli bo'laklar."""
    doc = Document(yol)
    natija = []
    for i, p in enumerate(_barcha_xatboshilar(doc)):
        matn = p.text
        if not matn.strip():
            continue
        bolaklar = []
        for qism in BELGI_NAQSHI.split(matn):
            if not qism:
                continue
            bolaklar.append({'belgi': bool(BELGI_NAQSHI.fullmatch(qism)), 'matn': qism})
        natija.append({'i': i, 'matn': matn, 'bolaklar': bolaklar,
                       'belgilar_soni': sum(1 for b in bolaklar if b['belgi'])})
    return natija


def _xatboshi_matnini_yoz(p, matn):
    """Xatboshi matnini almashtiradi, birinchi run formatlashini saqlab."""
    if p.runs:
        p.runs[0].text = matn
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(matn)


@login_required
def shablon_edit(request, turi):
    """Shablonni brauzerda tahrirlash."""
    _shablon_vakolati(request.user)
    if turi not in SHABLON_TURLARI:
        raise Http404
    fayl, nom, _belgilar = SHABLON_TURLARI[turi]
    joriy = os.path.join(SHABLONLAR, fayl)
    if not os.path.exists(joriy):
        messages.error(request, f'«{nom}» shabloni topilmadi.')
        return redirect('shablon_list')

    if request.method == 'POST':
        return _shablonni_saqla(request, turi, fayl, nom, joriy)

    return render(request, 'contracts/shablon_edit.html', {
        'turi': turi, 'nom': nom, 'fayl': fayl,
        'xatboshilar': _shablon_xatboshilari(joriy),
    })


def _shablonni_saqla(request, turi, fayl, nom, joriy):
    try:
        ozgarishlar = json.loads(request.POST.get('ozgarishlar') or '{}')
    except json.JSONDecodeError:
        messages.error(request, 'Ma\'lumot buzilib keldi, sahifani yangilab qayta urinib ko\'ring.')
        return redirect('shablon_edit', turi=turi)

    if not ozgarishlar:
        messages.info(request, 'Hech narsa o\'zgartirilmadi.')
        return redirect('shablon_edit', turi=turi)

    oldingi_belgilar = _belgilarni_sana(joriy)

    doc = Document(joriy)
    xatboshilar = list(_barcha_xatboshilar(doc))
    yozildi = 0
    for kalit, yangi in ozgarishlar.items():
        try:
            i = int(kalit)
        except (TypeError, ValueError):
            continue
        if not 0 <= i < len(xatboshilar):
            messages.error(request, 'Shablon tuzilishi o\'zgargan — sahifani yangilang.')
            return redirect('shablon_edit', turi=turi)
        _xatboshi_matnini_yoz(xatboshilar[i], yangi)
        yozildi += 1

    # Avval vaqtinchalik faylga yozib tekshiramiz
    vaqtinchalik = os.path.join(SHABLONLAR, f'.tahrir_{fayl}')
    doc.save(vaqtinchalik)
    xato = _shablonni_tekshir(vaqtinchalik, turi)
    if xato:
        os.remove(vaqtinchalik)
        messages.error(request, xato)
        return redirect('shablon_edit', turi=turi)

    zaxira_papka = os.path.join(SHABLONLAR, 'zaxira')
    os.makedirs(zaxira_papka, exist_ok=True)
    belgi = timezone.localtime().strftime('%Y%m%d-%H%M%S')
    shutil.copy2(joriy, os.path.join(zaxira_papka, f'{belgi}_{fayl}'))
    os.remove(joriy)
    os.replace(vaqtinchalik, joriy)

    amal_yoz(request.user, Amal.SHABLON_OZGARTIRDI, nom,
             f'{yozildi} ta qism o\'zgartirildi (onlayn)')
    messages.success(request, f'«{nom}» shabloni saqlandi — {yozildi} ta qism '
                              f'o\'zgartirildi. Eski nusxa zaxirada.')

    # Belgi kamayib qolgan bo'lsa jimgina o'tkazib yubormaymiz
    kamaygan = _kamaygan_belgilar(oldingi_belgilar, _belgilarni_sana(joriy))
    if kamaygan:
        messages.warning(
            request,
            'Diqqat: quyidagi belgilar kamaydi — ' + ', '.join(kamaygan) +
            '. Ular o\'chgan joyga mijoz ma\'lumoti yozilmaydi. '
            'Xato bo\'lsa zaxiradan qaytarish mumkin.')
    return redirect('shablon_edit', turi=turi)


def _belgilarni_sana(yol):
    """Shablondagi har bir belgining nechta uchraganini sanaydi."""
    doc = Document(yol)
    matn = '\n'.join(p.text for p in _barcha_xatboshilar(doc))
    hisob = {}
    for belgi in BELGI_NAQSHI.findall(matn):
        kalit = re.sub(r'\s+', ' ', belgi).strip()
        hisob[kalit] = hisob.get(kalit, 0) + 1
    return hisob


def _kamaygan_belgilar(oldin, keyin):
    kamaygan = []
    for belgi, soni in oldin.items():
        yangi = keyin.get(belgi, 0)
        if yangi < soni:
            kamaygan.append(f'{belgi} ({soni} → {yangi})')
    return kamaygan


# --------------------------------------------------------------- amallar tarixi

@login_required
def amallar(request):
    """Kim, qachon, nima qilgani — boshliqlar uchun."""
    _boss_required(request.user)
    qs = Amal.objects.select_related('kim')

    kim = request.GET.get('kim') or ''
    turi = request.GET.get('amal') or ''
    sanadan = request.GET.get('sanadan') or ''
    sanagacha = request.GET.get('sanagacha') or ''

    if kim:
        qs = qs.filter(kim_id=kim)
    if turi:
        qs = qs.filter(amal=turi)
    if sanadan:
        qs = qs.filter(vaqt__date__gte=sanadan)
    if sanagacha:
        qs = qs.filter(vaqt__date__lte=sanagacha)

    jami = qs.count()
    return render(request, 'contracts/amallar.html', {
        'amallar': qs[:300],
        'jami': jami,
        'foydalanuvchilar': User.objects.order_by('username'),
        'amal_turlari': Amal.AMAL_CHOICES,
        'tanlangan': {'kim': kim, 'amal': turi,
                      'sanadan': sanadan, 'sanagacha': sanagacha},
    })


# --------------------------------------------------------------- garov hujjati

def _garov_hujjati(request, pk, pdf=False):
    contract = get_object_or_404(Contract, pk=pk)
    if not _can_view_contract(request.user, contract):
        raise PermissionDenied('Bu shartnoma sizga tegishli emas.')
    if not garovi_bormi(contract):
        messages.info(request, 'Bu shartnomada garov hujjati tuzilmaydi '
                               '(ish haqi kafilligi).')
        return redirect('contract_detail', pk=pk)

    data = build_garov_docx(contract)
    if data is None:
        messages.error(request, 'Garov hujjati topilmadi.')
        return redirect('contract_detail', pk=pk)

    if pdf:
        try:
            data = docx_dan_pdf(data)
        except PdfImkoniYoq as xato:
            messages.error(request, str(xato))
            return redirect('contract_detail', pk=pk)
        tur = 'application/pdf'
        nom = f'garov_{contract.number}.pdf'
    else:
        tur = ('application/vnd.openxmlformats-officedocument'
               '.wordprocessingml.document')
        nom = f'garov_{contract.number}.docx'

    resp = HttpResponse(data, content_type=tur)
    resp['Content-Disposition'] = f'attachment; filename="{nom}"'
    return resp


@login_required
def contract_download_garov(request, pk):
    """Garov shartnomasi + baholash dalolatnomasi (Word)."""
    return _garov_hujjati(request, pk, pdf=False)


@login_required
def contract_download_garov_pdf(request, pk):
    """Garov shartnomasi + baholash dalolatnomasi (PDF)."""
    return _garov_hujjati(request, pk, pdf=True)
