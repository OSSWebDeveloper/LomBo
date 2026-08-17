# -*- coding: utf-8 -*-
"""Shartnoma to'plamining muqovasi (birinchi sahifa).

Namuna .doc fayli bo'lmagani uchun matn koddan yoziladi. Mijoz tanlagan
ko'rinish (2026-08-07) — ma'lumotlar ramkali quti ichida:

        Asia Invest Mikromoliya tashkiloti

           Рахмонова Шахноза Элмуродовна

               ──────────────────

        ┌──────────────────────┬───────────────────────────┐
        │ Кредит №             │ 159                       │
        │ Сана / Муддати       │ 05.08.2026 й — 04.08.2027 й│
        │ Микрокарз суммаси    │ 5 000 000,00              │
        │ Йиллик фоизи / Гаров │ 60% — Заргарлик буюмлари  │
        └──────────────────────┴───────────────────────────┘

              Бухоро шаҳри 2026 йил
"""
from django.conf import settings
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SHRIFT = 'Times New Roman'

SARLAVHA_PT = 30       # tashkilot nomi
ISM_PT = 26            # qarz oluvchi
MATN_PT = 12
BELGI_PT = 12          # quti ichidagi yozuvlar

KULRANG = '595959'     # bezak chizig'i uchun

PAGE_W, PAGE_H = 21.0, 29.7
MARGIN = 2.0


def _yoz(xatboshi, matn, *, olcham=MATN_PT, qalin=False, kursiv=False,
         tagchiziq=False, soya=False, oraliq=0, rang=None):
    """Matn bo'lagi. `soya` — Word'ning soya effekti, `oraliq` — harflar orasi (pt)."""
    run = xatboshi.add_run(matn)
    run.bold = qalin
    run.italic = kursiv
    run.underline = tagchiziq
    run.font.size = Pt(olcham)
    run.font.name = SHRIFT
    if rang:
        run.font.color.rgb = RGBColor.from_string(rang)
    if soya or oraliq:
        rPr = run._element.get_or_add_rPr()
        if soya:
            rPr.append(OxmlElement('w:shadow'))
        if oraliq:
            kenglik = OxmlElement('w:spacing')
            kenglik.set(qn('w:val'), str(int(oraliq * 20)))   # punktning 1/20 qismi
            rPr.append(kenglik)
    return run


def _bezak_chizigi(doc, en=6.0, keyin=6):
    """Sarlavha ostidagi qisqa, o'rtalangan chiziq."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(keyin)
    chet = (17.0 - en) / 2          # matn maydoni eni 17 sm
    par.paragraph_format.left_indent = Cm(chet)
    par.paragraph_format.right_indent = Cm(chet)
    _yoz(par, '', olcham=2)

    pBdr = OxmlElement('w:pBdr')
    past = OxmlElement('w:bottom')
    past.set(qn('w:val'), 'single')
    past.set(qn('w:sz'), '6')       # punktning 1/8 qismi
    past.set(qn('w:space'), '1')
    past.set(qn('w:color'), KULRANG)
    pBdr.append(past)
    par._p.get_or_add_pPr().append(pBdr)
    return par


def _p(doc, matn='', *, olcham=MATN_PT, qalin=False, kursiv=False,
       tagchiziq=False, markaz=True, oldin=0, keyin=6, soya=False, oraliq=0,
       rang=None):
    par = doc.add_paragraph()
    par.alignment = (WD_ALIGN_PARAGRAPH.CENTER if markaz
                     else WD_ALIGN_PARAGRAPH.LEFT)
    par.paragraph_format.space_before = Pt(oldin)
    par.paragraph_format.space_after = Pt(keyin)
    if matn:
        _yoz(par, matn, olcham=olcham, qalin=qalin, kursiv=kursiv,
             tagchiziq=tagchiziq, soya=soya, oraliq=oraliq, rang=rang)
    return par


def _bosh_joy(doc, sm):
    """Aniq balandlikdagi bo'sh oraliq (santimetrda).

    Bo'sh xatboshilar bilan emas, shu yo'l bilan qilinadi — asl muqovada
    bloklar sahifa bo'ylab ma'lum balandlikda turadi.
    """
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Cm(sm)
    run = par.add_run('')
    run.font.size = Pt(1)
    return par


def _jadval_chekinishi(jadval, sm):
    """Jadvalni chapdan chekintirish (python-docx'da tayyor xossasi yo'q)."""
    tblPr = jadval._tbl.tblPr
    ind = OxmlElement('w:tblInd')
    ind.set(qn('w:w'), str(int(sm * 567)))      # 1 sm = 567 twip
    ind.set(qn('w:type'), 'dxa')
    tblPr.append(ind)


def _jadval_chegarasi(jadval, qalinlik='6', rang='000000'):
    """Jadvalga chegara beradi.

    Uslub (`Table Grid`) ishlatilmaydi: muqova shartnoma hujjatining ichiga
    qo'yiladi, u yerda bu uslub yo'q — chegara ko'rinmay qolardi.
    """
    borders = OxmlElement('w:tblBorders')
    for tomon in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{tomon}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), qalinlik)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), rang)
        borders.append(el)
    jadval._tbl.tblPr.append(borders)


def _qatorlar_jadvali(doc, qatorlar, enlar, chekinish, olcham=BELGI_PT,
                      chegara=False, oraliq_pt=3):
    """Har qator — bir nechta katak, o'lchamlari qat'iy.

    `qatorlar` — [[(matn, qalin), ...], ...], `enlar` — ustun enlari (sm).
    """
    jadval = doc.add_table(rows=len(qatorlar), cols=len(enlar))
    jadval.autofit = False
    _jadval_chekinishi(jadval, chekinish)
    if chegara:
        _jadval_chegarasi(jadval)
    for i, qator in enumerate(qatorlar):
        for ustun, en in enumerate(enlar):
            katak = jadval.cell(i, ustun)
            katak.width = Cm(en)
            par = katak.paragraphs[0]
            par.paragraph_format.space_before = Pt(oraliq_pt)
            par.paragraph_format.space_after = Pt(oraliq_pt)
            if ustun < len(qator):
                matn, qalin = qator[ustun]
                _yoz(par, matn, olcham=olcham, qalin=qalin)
    return jadval


def _garov_matni(contract):
    """Muqovadagi «Гаров:» qatori — ta'minot turiga qarab."""
    if contract.collateral_type == contract.TYPE_ZARGARLIK:
        return 'Заргарлик буюмлари'
    if contract.collateral_type == contract.TYPE_TRANSPORT:
        return 'Транспорт воситаси'
    return 'Иш хаки кафиллиги'


def _summa_matni(contract):
    """«5 000 000,00» — muqovada summa tiyinlari bilan yoziladi."""
    from .shablondan import _raqam
    return f'{_raqam(contract.amount)},00'


def muqova_hujjati(contract):
    """Muqovani alohida Word hujjati sifatida qaytaradi."""
    org = settings.LOMBARD_ORG
    doc = Document()

    uslub = doc.styles['Normal']
    uslub.font.name = SHRIFT
    uslub.font.size = Pt(MATN_PT)

    for bolim in doc.sections:
        bolim.page_width = Cm(PAGE_W)
        bolim.page_height = Cm(PAGE_H)
        for chekka in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
            setattr(bolim, chekka, Cm(MARGIN))

    # Tashkilot nomi — qalin, kursivsiz, yengil soya bilan
    _bosh_joy(doc, 1.2)
    _p(doc, org['title_latin'], olcham=SARLAVHA_PT, qalin=True,
       soya=True, oraliq=1.0, keyin=0)

    # Qarz oluvchi ismi — muqovadagi ikkinchi yirik yozuv, qalin (2026-08-17)
    _bosh_joy(doc, 3.6)
    _p(doc, contract.borrower_fio, olcham=ISM_PT, qalin=True, keyin=0)

    _bosh_joy(doc, 0.5)
    _bezak_chizigi(doc, en=11.0, keyin=8)

    _bosh_joy(doc, 1.0)
    sana = contract.date.strftime('%d.%m.%Y')
    tugash = contract.end_date.strftime('%d.%m.%Y')
    _qatorlar_jadvali(doc, [
        [('Кредит №', True), (str(contract.number), True)],
        [('Сана / Муддати', True), (f'{sana} й — {tugash} й', True)],
        [('Микрокарз суммаси', True), (_summa_matni(contract), True)],
        [('Йиллик фоизи / Гаров', True),
         (f'{contract.interest_rate}% — {_garov_matni(contract)}', True)],
    ], enlar=[5.2, 6.5], chekinish=3.0, chegara=True)

    _bosh_joy(doc, 1.6)
    _p(doc, f'{org["city"]} {contract.date.year} йил',
       olcham=MATN_PT, qalin=True, kursiv=True, tagchiziq=True)

    return doc
