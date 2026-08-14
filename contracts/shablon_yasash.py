# -*- coding: utf-8 -*-
"""Asl .doc hujjatlaridan docxtpl shablonlarini yasaydi.

Asl fayldagi o'zgaruvchan qiymatlar (mijoz ismi, summa, sana...) o'rniga
`{{ nom }}` shaklidagi belgilar qo'yiladi. O'zgarmas matn va butun formatlash
asl faylning aynan o'zi bo'lib qoladi.

Har bir shablon uchta hujjatdan yig'iladi va hammasi bitta faylda bo'ladi:
    1) mikroqarz shartnomasi (+ garov shartnomasi, + baholash dalolatnomasi)
    2) ariza          — `m_ariza.docx`
    3) kredit qo'mitasi bayoni va farmoyish — `m_bayon.docx`

Manba fayllar:
    m159.docx  — yurist tuzatgan yangi shartnoma (zargarlik namunasi, 2026-08)
    m199.docx  — transport va kafillik namunasi (yurist tuzatishlari kodda
                 qo'lda qo'llanadi, chunki bu turlar uchun yangi namuna yo'q)
    m196.docx  — eski zargarlik namunasi, tarix uchun saqlanmoqda

Ishga tushirish:  python contracts/shablon_yasash.py
"""
import copy
import os
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# Skript sifatida ham (`python contracts/shablon_yasash.py`), paket ichidan ham
# ishlashi kerak — shuning uchun modul yo'li qo'lda qo'shiladi.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_ulash import hujjatni_ulash  # noqa: E402

PAPKA = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'shablonlar')

ZARGARLIK, TRANSPORT, KAFILLIK = 'zargarlik', 'transport', 'kafillik'


# --------------------------------------------------------------- yurish

def xatboshilar(doc):
    """Hujjatdagi barcha xatboshilar — jadval ichidagilari ham.

    XML daraxtining o'zidan yuriladi: birlashtirilgan (merged) kataklar
    sababli bir element bir necha marta qaytishining oldi olinadi.
    """
    for el in doc.element.body.iter(qn('w:p')):
        yield Paragraph(el, None)


def jadvallar(doc):
    """Barcha jadvallar — ichma-ichlari bilan, takrorlanmasdan."""
    for el in doc.element.body.iter(qn('w:tbl')):
        yield Table(el, None)


# --------------------------------------------------------------- almashtirish

def xatboshida_almashtir(p, eski, yangi, bir_marta=False):
    """Run'lar orasiga bo'lingan matnni ham almashtiradi.

    Topilgan joyning birinchi run'i formatlashi saqlanadi — ya'ni asl faylda
    qizil bo'lgan qism qizilligicha qoladi.

    `bir_marta=True` — yangi matn eskisini o'z ichiga olganda kerak
    (masalan matn oxiriga qo'shimcha yozilganda), aks holda almashtirish
    o'zini qayta-qayta topib cheksiz aylanib qoladi.
    """
    almashdi = 0
    while True:
        runs = p.runs
        toliq = ''.join(r.text for r in runs)
        # Asl faylda raqamlar ichida uzilmas bo'shliq (U+00A0) ishlatilgan.
        # Uzunligi bir xil bo'lgani uchun indekslar buzilmaydi.
        boshi = toliq.replace(' ', ' ').find(eski)
        if boshi < 0:
            return almashdi
        oxiri = boshi + len(eski)

        # Har bir run qaysi belgilarni egallaganini hisoblaymiz
        chegara = []
        joriy = 0
        for r in runs:
            chegara.append((joriy, joriy + len(r.text)))
            joriy += len(r.text)

        birinchi = None
        for i, (a, b) in enumerate(chegara):
            if b <= boshi or a >= oxiri:
                continue          # bu run tegishli emas
            kes_a = max(boshi, a) - a
            kes_b = min(oxiri, b) - a
            matn = runs[i].text
            if birinchi is None:
                birinchi = i
                runs[i].text = matn[:kes_a] + yangi + matn[kes_b:]
            else:
                runs[i].text = matn[:kes_a] + matn[kes_b:]
        almashdi += 1
        if bir_marta:
            return almashdi


def hujjatda_almashtir(doc, juftlar, bir_marta=False):
    """juftlar: [(eski_matn, yangi_matn), ...] — tartib muhim."""
    hisob = {eski: 0 for eski, _ in juftlar}
    for p in xatboshilar(doc):
        for eski, yangi in juftlar:
            hisob[eski] += xatboshida_almashtir(p, eski, yangi, bir_marta)
    return hisob


def xatboshida_topib_almashtir(doc, nishon, juftlar):
    """Faqat ichida `nishon` bo'lagi bor xatboshilarda almashtiradi.

    Global almashtirish ba'zi joyda yaramaydi: garov shartnomasida
    `{{ fio }}` goh qarz oluvchini, goh garovga qo'yuvchini bildiradi.
    """
    almashdi = 0
    for p in xatboshilar(doc):
        toliq = ''.join(r.text for r in p.runs).replace(' ', ' ')
        if nishon not in toliq:
            continue
        for eski, yangi in juftlar:
            almashdi += xatboshida_almashtir(p, eski, yangi, bir_marta=True)
    return almashdi


def naqsh_bilan_almashtir(doc, naqsh, yasovchi):
    """Regulyar ifoda topgan joyni `yasovchi(m)` qaytargan matnga almashtiradi.

    Bo'shliqlar soni aniq bo'lmagan joylarda (chiziqchalar, ustunlar orasi)
    ishlatiladi — topilgani baribir `xatboshida_almashtir` orqali yoziladi,
    ya'ni formatlash saqlanadi.
    """
    almashdi = 0
    for p in xatboshilar(doc):
        toliq = ''.join(r.text for r in p.runs).replace(' ', ' ')
        m = naqsh.search(toliq)
        if m:
            almashdi += xatboshida_almashtir(p, m.group(0), yasovchi(m), bir_marta=True)
    return almashdi


def natijani_chop(sarlavha, hisob):
    """Har bir juftlik nechta joyda ishlaganini ko'rsatadi."""
    print(f'  {sarlavha}:')
    for eski, n in hisob.items():
        belgi = 'OK ' if n else "YO'Q"
        print(f'    [{belgi}] {n:2} marta: {eski[:62]}')
    return all(hisob.values())


def qoldiqni_tekshir(doc, sozlar):
    """Namunadagi mijoz ma'lumoti qolib ketmaganini tekshiradi."""
    matn = '\n'.join(''.join(r.text for r in p.runs) for p in xatboshilar(doc))
    qolgan = [s for s in sozlar if s in matn]
    for s in qolgan:
        print(f"    [XATO] namuna qiymati qolib ketdi: {s}")
    return not qolgan


# --------------------------------------------------------------- telefon oralig'i

# 9-banddagi «Қарз олувчи» rekvizitlarining oxirgi qatori. Undan keyin darrov
# imzo jadvali boshlanadi — mijoz talabiga ko'ra (2026-08-14) telefon raqami
# bilan jadval orasida bo'sh joy qoladi.
TELEFON_REKVIZITI = 'Телефон: {{ telefon }}'


def telefondan_keyin_bosh_joy(doc):
    """Rekvizitlardagi telefon qatoridan keyin bo'sh xatboshi qo'shadi.

    Xatboshi nusxa olib yasaladi — shrift va oraliqlar o'sha qatorniki
    bo'lib qoladi, faqat matni olib tashlanadi.
    """
    qoshildi = 0
    for p in list(xatboshilar(doc)):
        if TELEFON_REKVIZITI not in ''.join(r.text for r in p.runs):
            continue
        bosh = copy.deepcopy(p._p)
        for run in bosh.findall(qn('w:r')):
            bosh.remove(run)
        p._p.addnext(bosh)
        qoshildi += 1
    return qoshildi


# --------------------------------------------------------------- zargarlik jadvali

def zargarlik_jadvalini_shablonla(doc):
    """Namuna qatorlarini bitta takrorlanuvchi qatorga aylantiradi."""
    # Aynan buyumlar jadvali: kamida 5 ustun va sarlavhasida «Кимматликлар»
    # bo'lgan qisqa katak (uni o'rab turgan katta jadvaldan farqlash uchun).
    nishon = None
    for t in jadvallar(doc):
        if len(t.columns) < 5:
            continue
        sarlavha = [c.text.strip() for c in t.rows[0].cells]
        if any('имматликлар' in s and len(s) < 40 for s in sarlavha):
            nishon = t
            break
    if nishon is None:
        return False, 'zargarlik jadvali topilmadi'

    qatorlar = nishon.rows
    if len(qatorlar) < 3:
        return False, f'jadvalda atigi {len(qatorlar)} qator bor'

    # 0 — sarlavha, 1..n-2 — buyumlar, oxirgisi — «Жами»
    birinchi_buyum = qatorlar[1]

    # Birinchi buyum qatorini namuna qatoriga aylantiramiz.
    # Bu qator docxtpl'ga tegishli emas — u render'dan keyin har bir buyum
    # uchun nusxalanadi (shablondan.py dagi _buyumlar_jadvali).
    qiymatlar = ['#IDX#', '#NOMI#', '#SONI# та', '#OGIR# гр', '#PROBA#', '#SUMMA#']
    kataklar = birinchi_buyum.cells
    for i, katak in enumerate(kataklar):
        if i >= len(qiymatlar):
            break
        p = katak.paragraphs[0]
        if p.runs:
            p.runs[0].text = qiymatlar[i]
            for r in p.runs[1:]:
                r.text = ''
        else:
            p.add_run(qiymatlar[i])
        for qoshimcha in katak.paragraphs[1:]:
            for r in qoshimcha.runs:
                r.text = ''
        # Asl faylda kataklarda 1 sm manfiy chekinish bor, matn esa oldiga
        # bo'shliq qo'yilib to'g'rilangan. Biz katak matnini butunlay qayta
        # yozamiz — bo'shliqlar yo'qolgani uchun chekinishni ham nolga
        # tushiramiz, aks holda raqam katakdan chiqib ketib ko'rinmay qoladi.
        for qism in katak.paragraphs:
            qism.paragraph_format.left_indent = 0

    # Ortiqcha namuna qatorlarini o'chiramiz («Жами» qatoridan tashqari)
    ortiqcha = list(qatorlar[2:-1])
    for qator in ortiqcha:
        qator._element.getparent().remove(qator._element)

    return True, f'jadval shablonlandi, {len(ortiqcha)} ta ortiqcha qator olindi'


# =============================================================== ARIZA

# Namunadagi mijoz: Рахмонова Шахноза Элмуродовна, 5 000 000 сўм, 12 ой, 60%
ARIZA_TAMINOT = ('Рахмонова Шахноза Элмуродовна (узимга) тегишли заргарлик '
                 'буюмларини гаровга такдим этаман')
ARIZA_TAMINOT_YANGI = {
    # Zargarlikda jumla garovga qo'yuvchiga bog'liq: qarz oluvchining o'zi
    # bo'lsa «узимга тегишли...», boshqa shaxs bo'lsa uning ismi yoziladi.
    ZARGARLIK: '{{ ariza_taminot }}',
    TRANSPORT: '{{ garov_mulki }}ни гаровга такдим этаман',
    KAFILLIK: ('{{ kafil }}нинг {{ kafillik_summa }} сўмлик иш хакки '
               'кафиллигини такдим этаман'),
}

ARIZA_UMUMIY = [
    # Uzunroq va aniqroqlari birinchi
    ('5 000 000 (беш миллион)', '{{ summa }}'),
    ('ойида ўртача 5 000 000 сўм', 'ойида ўртача {{ daromad }} сўм'),
    ('12 ой муддатга', '{{ muddat }} ой муддатга'),
    ('йилига 60 фоиз', 'йилига {{ foiz }} фоиз'),
    # Viloyat va IIV bo'lim raqami bitta belgida: yashil biometrik pasportda
    # bo'lim raqami bo'lmaydi, faqat viloyat qoladi (qarang: _umumiy).
    ('09.02.2023-йилда , Бухоро вилояти 6224 - сонли ИИВ томонидан берилган',
     '{{ pasport_sana }}-йилда, {{ pasport_bergan }} ИИВ томонидан берилган'),
    ('№ 2540542', '№ {{ pasport_soni }}'),
    ('АD', '{{ pasport_seriya }}'),
    ('Менинг доимий яшаш манзилим: Бухоро вил,Гиждувон туман, Чогдаре МФЙ, '
     'Чогдаре кишлоги.', 'Менинг доимий яшаш манзилим: {{ manzil }}'),
    ('05 август 2026 йил', '{{ sana_soz }} йил'),
    ('Рахмонова Шахноза Элмуродовна', '{{ fio }}'),
]

# «Телефон ракам  1)____ 2) ____ 3) ____» — chiziqchalar soni aniq bo'lmagani
# uchun naqsh bilan topiladi. Uchala o'rin ham to'ldiriladi.
TELEFON_NAQSHI = re.compile(
    r'(Телефон ракам\s*1\))\s*_+(\s*2\))\s*_+(\s*3\))\s*_+')

# «Менинг иш жойим ва унинг манзили: ______»
ISH_JOYI_NAQSHI = re.compile(r'(Менинг иш жойим ва унинг манзили:)\s*_+')


def ariza_hujjati(tur):
    doc = Document(os.path.join(PAPKA, 'm_ariza.docx'))
    juftlar = [(ARIZA_TAMINOT, ARIZA_TAMINOT_YANGI[tur])] + ARIZA_UMUMIY
    ok = natijani_chop('ariza', hujjatda_almashtir(doc, juftlar))

    n = naqsh_bilan_almashtir(
        doc, TELEFON_NAQSHI,
        lambda m: (f'{m.group(1)} {{{{ telefon }}}}'
                   f'{m.group(2)} {{{{ telefon2 }}}}'
                   f'{m.group(3)} {{{{ telefon3 }}}}'))
    print(f"    [{'OK ' if n else 'YO`Q'}] {n:2} marta: Телефон ракам 1) 2) 3)")
    ok = ok and bool(n)

    m = naqsh_bilan_almashtir(doc, ISH_JOYI_NAQSHI,
                              lambda x: f'{x.group(1)} {{{{ ish_joyi }}}}')
    print(f"    [{'OK ' if m else 'YO`Q'}] {m:2} marta: Менинг иш жойим...")
    return doc, ok and bool(m)


# =============================================================== BAYON

BAYON_TAMINOT = {
    # Garovga qo'yuvchi qarz oluvchining o'zi bo'lmasligi mumkin, shuning uchun
    # «узига тегишли» iborasi kontekstdan keladi (qarang: zargarlik_konteksti).
    ZARGARLIK: [
        ('Гаров таьминоти сифатида фукаро Рахмонова Шахноза Элмуродовнага '
         'тегишли заргарлик буюмлари қабул қилинсин.',
         'Гаров таьминоти сифатида фукаро {{ garov_mulki }} қабул қилинсин.'),
        ('накд пулда, Рахмонова Шахноза Элмуродовнага тегишли заргарлик '
         'буюмлари гарови асосида',
         'накд пулда, {{ garov_mulki }} гарови асосида'),
        ('Рахмонова Шахноза Элмуродовна узига тегишли заргарлик буюмларини '
         'гаровга куйилишини', '{{ garov_mulki_egalik }}ни гаровга куйилишини'),
    ],
    TRANSPORT: [
        ('Гаров таьминоти сифатида фукаро Рахмонова Шахноза Элмуродовнага '
         'тегишли заргарлик буюмлари қабул қилинсин.',
         'Гаров таьминоти сифатида {{ garov_mulki }} қабул қилинсин.'),
        ('накд пулда, Рахмонова Шахноза Элмуродовнага тегишли заргарлик '
         'буюмлари гарови асосида',
         'накд пулда, {{ garov_mulki }} гарови асосида'),
        ('Рахмонова Шахноза Элмуродовна узига тегишли заргарлик буюмларини '
         'гаровга куйилишини', '{{ garov_mulki }} гаровга куйилишини'),
    ],
    KAFILLIK: [
        ('гаров таьминоти сифатида Рахмонова Шахноза Элмуродовна узига тегишли '
         'заргарлик буюмларини гаровга куйилишини',
         'таъминот сифатида {{ kafil }}нинг {{ kafillik_summa }} сўмлик иш хакки '
         'кафиллиги такдим этилишини'),
        ('таъминот сифатида Рахмонова Шахноза Элмуродовна узига тегишли '
         'заргарлик буюмларини гаровга куйилишини',
         'таъминот сифатида {{ kafil }}нинг {{ kafillik_summa }} сўмлик иш хакки '
         'кафиллиги такдим этилганини'),
        ('Гаров таьминоти сифатида фукаро Рахмонова Шахноза Элмуродовнага '
         'тегишли заргарлик буюмлари қабул қилинсин.',
         'Таъминот сифатида {{ kafil }}нинг {{ kafillik_summa }} сўмлик иш хакки '
         'кафиллиги қабул қилинсин.'),
        ('накд пулда, Рахмонова Шахноза Элмуродовнага тегишли заргарлик '
         'буюмлари гарови асосида',
         'накд пулда, {{ kafil }}нинг {{ kafillik_summa }} сўмлик иш хакки '
         'кафиллиги асосида'),
    ],
}

BAYON_UMUMIY = [
    ('№159-сонли', '№{{ raqam }}-сонли'),
    ('№159', '№{{ raqam }}'),
    ('05.08.2026', '{{ sana }}'),
    ('5 000 000 (беш миллион)', '{{ summa }}'),
    ('12 ой муддатга', '{{ muddat }} ой муддатга'),
    ('йиллик 60 фоиз', 'йиллик {{ foiz }} фоиз'),
    ('Рахмонова Шахноза Элмуродовнага', '{{ fio }}га'),
    ('Рахмонова Шахноза Элмуродовна', '{{ fio }}'),
]


def bayon_hujjati(tur):
    doc = Document(os.path.join(PAPKA, 'm_bayon.docx'))
    juftlar = BAYON_TAMINOT[tur] + BAYON_UMUMIY
    ok = natijani_chop('bayon', hujjatda_almashtir(doc, juftlar))
    return doc, ok


# =============================================================== zargarlik (m159)

ZARGARLIK_JUFTLAR = [
    # Manzil + telefon — 9-banddagi «Қарз олувчи» rekvizitlari
    ('Манзил: Бухоро вилояти,Гиждувон туман, Чогдаре МФЙ, Чогдаре кишлоги.',
     'Манзил: {{ manzil }}. Телефон: {{ telefon }}'),
    # Pasport asl faylda to'rt xil yozilgan (chiziqcha bor/yo'q, uzun tire)
    ('Бухоро вилояти 6224 сонли ИИВ томонидан 09.02.2023-йилда берилган '
     'АD №2540542 ракамли шахс гувохномаси', '{{ pasport }}'),
    ('Бухоро вилояти 6224 - сонли ИИВ томонидан 09.02.2023-йилда берилган '
     'АD №2540542 ракамли шахс гувохномаси', '{{ pasport }}'),
    ('Бухоро вилояти 6224 – сонли ИИВ томонидан 09.02.2023-йилда берилган '
     'АD №2540542 ракамли шахс гувохномаси', '{{ pasport }}'),
    ('Бухоро вилояти 6224 томонидан 09.02.2023-йилда берилган '
     'АD №2540542 ракамли шахс гувохномаси', '{{ pasport }}'),
    ('Бухоро вилояти, Гиждувон тумани, Чогдаре МФЙ, Чогдаре кишлоги', '{{ manzil }}'),
    ('Гиждувон туман, Чогдаре МФЙ, Чогдаре кишлоги', '{{ manzil }}'),
    ('6 000 000 (олти миллион)', '{{ garov_baho }}'),
    ('5 000 000 (беш миллион)', '{{ summa }}'),
    ('Микрокарз шартномаси №159', 'Микрокарз шартномаси №{{ raqam }}'),
    ('Гаров шартнома 159', 'Гаров шартнома {{ garov_raqam }}'),
    ('далолатномаси №159', 'далолатномаси №{{ raqam }}'),
    ('№159-сонли', '№{{ raqam }}-сонли'),
    ('Рахмонова Шахноза Элмуродовнага', '{{ fio }}га'),
    ('Рахмонова Шахноза Элмуродовна', '{{ fio }}'),
    ('04.08.2027', '{{ tugash }}'),
    ('05.08.2026', '{{ sana }}'),
    ('12 (ун икки)', '{{ muddat }} ({{ muddat_soz }})'),
    ('60% (олтмиш)', '{{ foiz }}% ({{ foiz_soz }})'),
    # Jadvalning «Жами» qatori
    ('2 та', '{{ jami_soni }} та'),
    ('6,70 гр', '{{ jami_ogirligi }} гр'),
    ('6 000 000', '{{ garov_baho_raqam }}'),
]


# Garov shartnomasi va baholash dalolatnomasida «гаровга қўювчи» qarz
# oluvchining o'zi bo'lmasligi mumkin. Global almashtirishdan keyin aynan shu
# joylar alohida belgilarga bog'lanadi — qolgan `{{ fio }}` lar (masalan
# «микрокарз олувчи фукаро {{ fio }}») qarz oluvchida qoladi.
GAROV_BERUVCHI_JOYLARI = [
    # (xatboshini tanish uchun bo'lak, [(eski, yangi), ...])
    ('хамда гаровга кўювчи:',
     [('{{ fio }} ({{ pasport }})', '{{ garov_fio }} ({{ garov_pasport }})')]),
    ('келишув далолатномасига асосан гаровга куйиладиган',
     [('{{ fio }}га тегишли', '{{ garov_fio }}га тегишли')]),
    ('Кредит таъминоти сифатида гаровга кўйиладиган',
     [('{{ fio }}га тегишли', '{{ garov_fio }}га тегишли')]),
    ('Биз, қуйида имзо чекувчилар',
     [('қарз олувчи ва гаровга куювчи {{ fio }} ({{ pasport }} ,Манзил: {{ manzil }})',
       '{{ dalolatnoma_taraflar }}')]),
    ('Гаровга куювчи:',                       # dalolatnoma imzo qatori
     [('{{ fio }}', '{{ garov_fio }}')]),
]

# 8-banddagi «Гаровга қўювчи:» rekvizit bloki — sarlavhadan keyingi qatorlar
GAROV_REKVIZIT_SARLAVHASI = 'Гаровга қўювчи:'
GAROV_REKVIZIT_JUFTLARI = [
    ('Манзил: {{ manzil }}', 'Манзил: {{ garov_manzil }}'),
    ('{{ pasport }}', '{{ garov_pasport }}'),
    ('{{ fio }}', '{{ garov_fio }}'),
]


def garov_beruvchini_shablonla(doc):
    """Garovga qo'yuvchiga tegishli joylarni alohida belgilarga bog'laydi."""
    ok = True
    for nishon, juftlar in GAROV_BERUVCHI_JOYLARI:
        n = xatboshida_topib_almashtir(doc, nishon, juftlar)
        print(f"    [{'OK ' if n else 'YO`Q'}] {n:2} marta: {nishon[:52]}")
        ok = ok and bool(n)

    # Rekvizit blokida qatorlar alohida xatboshida turadi va ularda «гаровга
    # қўювчи» so'zi yo'q — shuning uchun sarlavhadan keyingi qatorlar olinadi.
    pars = list(xatboshilar(doc))
    topildi = 0
    for i, p in enumerate(pars):
        matn = ' '.join(''.join(r.text for r in p.runs).split())
        if matn != GAROV_REKVIZIT_SARLAVHASI:
            continue
        topildi += 1
        for keyingi in pars[i + 1:i + 4]:
            for eski, yangi in GAROV_REKVIZIT_JUFTLARI:
                xatboshida_almashtir(keyingi, eski, yangi, bir_marta=True)
    print(f"    [{'OK ' if topildi else 'YO`Q'}] {topildi:2} marta: "
          f'{GAROV_REKVIZIT_SARLAVHASI} rekvizit bloki')
    return ok and bool(topildi)


def zargarlik_shabloni():
    doc = Document(os.path.join(PAPKA, 'm159.docx'))

    ok, xabar = zargarlik_jadvalini_shablonla(doc)
    print(f'  jadval: {xabar}')

    ok &= natijani_chop('shartnoma', hujjatda_almashtir(doc, ZARGARLIK_JUFTLAR))
    print('  garovga qo‘yuvchi:')
    ok &= garov_beruvchini_shablonla(doc)
    return _yakunla(doc, ZARGARLIK, 'zargarlik.docx', ok,
                    ['Рахмонова', 'Чогдаре', '2540542', '6224'])


# =============================================================== transport / kafillik (m199)

# 199-faylning mikroqarz qismidagi umumiy qiymatlar
Q199_UMUMIY = [
    ('61013-сонли ИИВ томонидан 20.05.2025-йилда берилган АE№2751664 '
     'ракамли шахс гувохномаси', '{{ pasport }}'),
    # 9-banddagi rekvizitlar — bu yerga telefon ham qo'shiladi
    ('Манзил:Бухоро вилояти, Бухоро  туман, Работикалмок МФЙ,Тикончи кўчаси.',
     'Манзил: {{ manzil }}. Телефон: {{ telefon }}'),
    # Kirish xatboshisida manzil boshqacha yozilgan («кучаси»)
    ('Манзил:Бухоро вилояти, Бухоро  туман, Работикалмок МФЙ,Тикончи кучаси',
     'Манзил: {{ manzil }}'),
    ('7 000 000 (етти миллион)', '{{ summa }}'),
    ('Микрокарз шартномаси №199', 'Микрокарз шартномаси №{{ raqam }}'),
    ('Содиков Ботир Нусратович', '{{ fio }}'),
    ('30.10.2026', '{{ tugash }}'),
    ('31.10.2025', '{{ sana }}'),
    ('12(ўн икки)', '{{ muddat }}({{ muddat_soz }})'),
    ('60% (олтмиш )', '{{ foiz }}% ({{ foiz_soz }})'),
]

# Yurist 2026-08 da kiritgan tuzatishlar — m159 da ular allaqachon bor,
# m199 ga esa shu yerda qo'llanadi.
YURIST_TUZATISHLARI = [
    ('мувофиқ Суд тартибида ҳал қилинадилар',
     'мувофиқ Нотариал идоранинг Ижро хати хамда Суд тартибида ҳал қилинадилар'),
    ('Телефон: (99891) 415-00-87;', 'Телефон: 55 310 00 87, 91 415-00-87;'),
]

# Garov shartnomasidagi bank rekvizitlari o'zgargan (faqat garov qismida).
# Bo'shliqlar soni aniq bo'lmagani uchun naqsh bilan almashtiriladi.
GAROV_NAQSHLARI = [
    (re.compile(r'(Х/Р\s+)20216000005489627001'),
     lambda m: m.group(1) + '20216000405489627001'),
    (re.compile(r'(М\.Ф\.О\.\s+)01095'), lambda m: m.group(1) + '01137'),
    # Dalolatnoma sarlavhasiga shartnoma raqami qo'shildi
    (re.compile(r'келишуви\s+асосида\s+бахолаш\s+далолатномаси\s*$'),
     lambda m: m.group(0).rstrip() + ' №{{ raqam }}'),
]

# 1.1-banddagi ta'minot jumlasi — turga qarab boshqacha
KAFILLIK_TAMINOT = (
    'Жураев Азизбек Носировичнинг  10 000 000 (ўн миллион) сўмлик '
    'иш хакки кафиллиги  такдим килинади.')
KAFILLIK_YANGI = '{{ kafil }}нинг {{ kafillik_summa }} сўмлик иш хакки кафиллиги такдим килинади.'
TRANSPORT_YANGI = ('{{ garov_egasi }}га тегишли, давлат раками {{ davlat_raqami }} бўлган, '
                   '{{ rusumi }} русумли транспорт воситаси гаровга қўйилади.')

# Garov shartnomasi va dalolatnomadagi transport qiymatlari
TRANSPORT_JUFTLAR = [
    ('“Express Alligator Bukhara” МЧЖга тегишли, давлат раками 80 9511АA бўлган, '
     'KRONE SDR27 русумли YARIM TIRKAMA REFRIJERATOR', '{{ garov_mulki }}'),
    ('“Express Alligator Bukhara” МЧЖга тегишли, давлат раками 80 9511 АA бўлган, '
     'KRONE SDR27 русумли YARIM TIRKAMA REFRIJERATOR', '{{ garov_mulki }}'),
    ('70 000 000 (етмиш миллион)', '{{ garov_baho }}'),
    ('50 000 000(эллик миллион)', '{{ summa_raqam_soz }}'),
    ('Гаров шартнома 34', 'Гаров шартнома {{ garov_raqam }}'),
    ('№34-сонли', '№{{ raqam }}-сонли'),
    ('04.04.2025', '{{ sana }}'),
    ('Бахшиллоева Нозима Бахтиёровна', '{{ fio }}'),
    ('Бахшиллоева Дилноза Бахтиёровна', '{{ garov_rahbari }}'),
    ('Д.Б.Бахшиллоева', '{{ garov_rahbari_qisqa }}'),   # imzo qatoridagi qisqartma
    ('“Express Alligator Bukhara” МЧЖ', '{{ garov_egasi }}'),
    ('Бухоро шахар 6206-сонли ИИВ томонидан 08.04.2021 йилда берилган '
     'АD 0301390 ракамли шахсни тасдикловчи хужжат', '{{ pasport }}'),
    ('Бухоро шахар, Шайхон кучаси, 156 – уй', '{{ manzil }}'),
    ('80 9511 АА', '{{ davlat_raqami }}'),
    ('KRONE SDR27', '{{ rusumi }}'),
    ('QORA CHERNIY', '{{ rangi }}'),
    ('WKESD000000713326', '{{ shassi }}'),
    ('2016-yil', '{{ yili }}'),
    ('AAG 0949387 / 14.02.2023 йил', '{{ texpasport }}'),
]


def _garov_qismini_ochir(doc):
    """Kafillik uchun garov shartnomasi va dalolatnomani olib tashlaydi."""
    body = doc.element.body
    ochirildi = 0
    for el in list(body):
        if el.tag != qn('w:tbl'):
            continue
        matn = ''.join(t.text or '' for t in el.iter(qn('w:t')))
        if 'аров шартнома' in matn and 'ШАРТНОМА  ПРЕДМЕТИ' in matn:
            body.remove(el)
            ochirildi += 1
    return ochirildi


def _yakunla(doc, tur, nomi, ok, qoldiq_sozlari):
    """Arizani va bayonni ulab, shablonni saqlaydi."""
    n = telefondan_keyin_bosh_joy(doc)
    print(f"    [{'OK ' if n else 'YO`Q'}] {n:2} marta: telefondan keyin bo'sh qator")
    ok = ok and bool(n)

    ariza, ariza_ok = ariza_hujjati(tur)
    bayon, bayon_ok = bayon_hujjati(tur)
    hujjatni_ulash(doc, ariza)
    hujjatni_ulash(doc, bayon)

    toza = qoldiqni_tekshir(doc, qoldiq_sozlari)
    chiqish = os.path.join(PAPKA, nomi)
    doc.save(chiqish)
    print(f'  -> {chiqish}')
    return bool(ok and ariza_ok and bayon_ok and toza)


def _m199_shabloni(nomi, tur, juftlar, taminot_yangi, garovni_ochir=False):
    doc = Document(os.path.join(PAPKA, 'm199.docx'))

    if garovni_ochir:
        n = _garov_qismini_ochir(doc)
        print(f'  garov qismi olib tashlandi: {n} ta blok')

    juftlar = [(KAFILLIK_TAMINOT, taminot_yangi)] + juftlar + YURIST_TUZATISHLARI
    ok = natijani_chop('shartnoma', hujjatda_almashtir(doc, juftlar))

    if not garovni_ochir:
        for naqsh, yasovchi in GAROV_NAQSHLARI:
            n = naqsh_bilan_almashtir(doc, naqsh, yasovchi)
            print(f"    [{'OK ' if n else 'YO`Q'}] {n:2} marta: {naqsh.pattern[:52]}")
            ok = ok and bool(n)

    qoldiq = ['Содиков', 'Работикалмок', '2751664']
    if not garovni_ochir:
        qoldiq += ['Бахшиллоева', 'Express Alligator', 'KRONE']
    return _yakunla(doc, tur, nomi, ok, qoldiq)


def kafillik_shabloni():
    return _m199_shabloni('kafillik.docx', KAFILLIK, list(Q199_UMUMIY),
                          KAFILLIK_YANGI, garovni_ochir=True)


def transport_shabloni():
    return _m199_shabloni('transport.docx', TRANSPORT,
                          list(Q199_UMUMIY) + TRANSPORT_JUFTLAR, TRANSPORT_YANGI)


if __name__ == '__main__':
    natijalar = []
    print('ZARGARLIK shabloni:')
    natijalar.append(zargarlik_shabloni())
    print('\nKAFILLIK shabloni:')
    natijalar.append(kafillik_shabloni())
    print('\nTRANSPORT shabloni:')
    natijalar.append(transport_shabloni())
    print('\nNatija:', 'hammasi joyida' if all(natijalar) else 'XATOLAR BOR')
    sys.exit(0 if all(natijalar) else 1)
